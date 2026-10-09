from __future__ import annotations

import asyncio
import hashlib
import io
import posixpath
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Awaitable, Callable
from xml.etree import ElementTree

from docx import Document
from docx.document import Document as DocumentObject


class DocumentEngineError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class DocumentVerification:
    valid: bool
    format: str
    byte_size: int
    checksum: str
    page_count: int | None


@dataclass(frozen=True)
class DocumentExportResult:
    document: bytes
    backend: str
    quality: str
    warnings: tuple[str, ...] = ()


ResolveExecutable = Callable[[], str | None]
RunLibreOffice = Callable[[str, Path, Path, Path], Awaitable[None]]
RunCupsfilter = Callable[[str, Path, Path], Awaitable[None]]

_FIXED_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
_FIXED_CORE_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)
_RELATIONSHIP_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
_WORD_NAMESPACE = (
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
)


class DocumentEngine:
    def __init__(
        self,
        *,
        resolve_executable: ResolveExecutable | None = None,
        resolve_cupsfilter: ResolveExecutable | None = None,
        run: RunLibreOffice | None = None,
        run_cupsfilter: RunCupsfilter | None = None,
    ) -> None:
        self._resolve_executable = resolve_executable or _find_libreoffice
        self._resolve_cupsfilter = resolve_cupsfilter or _find_cupsfilter
        self._run = run or _run_libreoffice
        self._run_cupsfilter = run_cupsfilter or _run_cupsfilter

    def create_docx(self, content: dict[str, object]) -> bytes:
        title = _required_text(content.get("title"))
        blocks = content.get("blocks")
        if not isinstance(blocks, list) or not blocks:
            raise DocumentEngineError(
                "document_content_invalid",
                "Document blocks are required",
            )

        document = Document()
        _set_deterministic_properties(document, title)
        document.add_paragraph(title, style="Title")
        for block in blocks:
            _append_block(document, block)

        output = io.BytesIO()
        document.save(output)
        normalized = _normalize_ooxml(output.getvalue())
        self.verify(document=normalized, format_="docx")
        return normalized

    async def export_pdf(self, document: bytes) -> DocumentExportResult:
        self.verify(document=document, format_="docx")
        executable = self._resolve_executable()
        if not executable:
            return await self._export_pdf_with_text_fallback(document)

        with tempfile.TemporaryDirectory(
            prefix="realmflow-document-export-"
        ) as directory:
            root = Path(directory)
            source_path = root / "source.docx"
            output_directory = root / "output"
            output_path = output_directory / "source.pdf"
            profile_path = root / "profile"
            output_directory.mkdir()
            profile_path.mkdir()
            source_path.write_bytes(document)
            try:
                await self._run(
                    executable,
                    source_path,
                    output_path,
                    profile_path,
                )
                converted = output_path.read_bytes()
            except DocumentEngineError:
                raise
            except Exception as error:
                raise DocumentEngineError(
                    "document_export_failed",
                    "PDF export failed",
                ) from error
            self.verify(document=converted, format_="pdf")
            return DocumentExportResult(
                document=converted,
                backend="libreoffice",
                quality="print",
            )

    async def _export_pdf_with_text_fallback(
        self,
        document: bytes,
    ) -> DocumentExportResult:
        executable = self._resolve_cupsfilter()
        if not executable:
            raise DocumentEngineError(
                "document_export_unavailable",
                "No local PDF export backend is available",
            )
        text = _extract_docx_plain_text(document)
        if not text.strip():
            raise DocumentEngineError(
                "document_export_failed",
                "DOCX text fallback found no exportable text",
            )

        with tempfile.TemporaryDirectory(
            prefix="realmflow-document-export-"
        ) as directory:
            root = Path(directory)
            source_path = root / "source.txt"
            output_path = root / "source.pdf"
            source_path.write_text(text, encoding="utf-8")
            try:
                await self._run_cupsfilter(
                    executable,
                    source_path,
                    output_path,
                )
                converted = output_path.read_bytes()
            except DocumentEngineError:
                raise
            except Exception as error:
                raise DocumentEngineError(
                    "document_export_failed",
                    "Text PDF export failed",
                ) from error
            self.verify(document=converted, format_="pdf")
            return DocumentExportResult(
                document=converted,
                backend="cupsfilter_text",
                quality="degraded_text",
                warnings=(
                    "PDF was generated from extracted plain text; original DOCX layout was not preserved.",
                ),
            )

    def verify(
        self,
        *,
        document: bytes,
        format_: str,
    ) -> DocumentVerification:
        try:
            if format_ == "docx":
                page_count = _verify_docx(document)
            elif format_ == "pdf":
                page_count = _verify_pdf(document)
            else:
                raise DocumentEngineError(
                    "document_format_unsupported",
                    "Document format is unsupported",
                )
        except DocumentEngineError:
            raise
        except Exception as error:
            raise DocumentEngineError(
                "document_verification_failed",
                f"{format_.upper()} verification failed",
            ) from error
        return DocumentVerification(
            valid=True,
            format=format_,
            byte_size=len(document),
            checksum=f"sha256:{hashlib.sha256(document).hexdigest()}",
            page_count=page_count,
        )


def libreoffice_pdf_command(
    *,
    executable: str,
    source_path: Path,
    output_path: Path,
    profile_path: Path,
) -> tuple[str, ...]:
    return (
        executable,
        "--headless",
        "--nologo",
        "--nodefault",
        "--nofirststartwizard",
        f"-env:UserInstallation={profile_path.as_uri()}",
        "--convert-to",
        "pdf:writer_pdf_Export",
        "--outdir",
        str(output_path.parent),
        str(source_path),
    )


def _find_libreoffice() -> str | None:
    discovered = shutil.which("libreoffice") or shutil.which("soffice")
    if discovered:
        return discovered
    macos_path = Path(
        "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    )
    return str(macos_path) if macos_path.is_file() else None


def _find_cupsfilter() -> str | None:
    discovered = shutil.which("cupsfilter")
    if discovered:
        return discovered
    macos_path = Path("/usr/sbin/cupsfilter")
    return str(macos_path) if macos_path.is_file() else None


async def _run_libreoffice(
    executable: str,
    source_path: Path,
    output_path: Path,
    profile_path: Path,
) -> None:
    process = await asyncio.create_subprocess_exec(
        *libreoffice_pdf_command(
            executable=executable,
            source_path=source_path,
            output_path=output_path,
            profile_path=profile_path,
        ),
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )
    try:
        stdout, stderr = await asyncio.wait_for(
            process.communicate(),
            timeout=120,
        )
    except TimeoutError as error:
        process.kill()
        await process.wait()
        raise DocumentEngineError(
            "document_export_failed",
            "PDF export timed out",
        ) from error
    if process.returncode != 0 or not output_path.is_file():
        detail = (stderr or stdout).decode("utf-8", errors="replace").strip()
        raise DocumentEngineError(
            "document_export_failed",
            detail[:500] or "PDF export failed",
        )


async def _run_cupsfilter(
    executable: str,
    source_path: Path,
    output_path: Path,
) -> None:
    process = await asyncio.create_subprocess_exec(
        executable,
        "-i",
        "text/plain",
        "-m",
        "application/pdf",
        str(source_path),
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )
    try:
        stdout, stderr = await asyncio.wait_for(
            process.communicate(),
            timeout=120,
        )
    except TimeoutError as error:
        process.kill()
        await process.wait()
        raise DocumentEngineError(
            "document_export_failed",
            "Text PDF export timed out",
        ) from error
    if process.returncode != 0 or not stdout:
        detail = stderr.decode("utf-8", errors="replace").strip()
        raise DocumentEngineError(
            "document_export_failed",
            detail[:500] or "Text PDF export failed",
        )
    output_path.write_bytes(stdout)


def _extract_docx_plain_text(document: bytes) -> str:
    try:
        parsed = Document(io.BytesIO(document))
    except Exception as error:
        raise DocumentEngineError(
            "document_export_failed",
            "DOCX text fallback could not read the document",
        ) from error

    lines: list[str] = []
    for paragraph in parsed.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in parsed.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n".join(lines) + "\n"


def _set_deterministic_properties(
    document: DocumentObject,
    title: str,
) -> None:
    properties = document.core_properties
    properties.title = title
    properties.author = "RealmFlow"
    properties.last_modified_by = "RealmFlow"
    properties.created = _FIXED_CORE_TIME
    properties.modified = _FIXED_CORE_TIME
    properties.revision = 1


def _append_block(document: DocumentObject, value: object) -> None:
    if not isinstance(value, dict):
        _invalid_content()
    kind = value.get("kind")
    if kind == "heading":
        level = value.get("level")
        text = _required_text(value.get("text"))
        if not isinstance(level, int) or isinstance(level, bool) or not 1 <= level <= 9:
            _invalid_content()
        document.add_heading(text, level=level)
        return
    if kind == "paragraph":
        document.add_paragraph(_required_text(value.get("text")))
        return
    if kind == "bullet_list":
        items = value.get("items")
        if not isinstance(items, list) or not items:
            _invalid_content()
        for item in items:
            document.add_paragraph(_required_text(item), style="List Bullet")
        return
    if kind == "table":
        rows = value.get("rows")
        if (
            not isinstance(rows, list)
            or not rows
            or not all(isinstance(row, list) and row for row in rows)
        ):
            _invalid_content()
        width = len(rows[0])
        if any(len(row) != width for row in rows):
            _invalid_content()
        table = document.add_table(rows=len(rows), cols=width)
        for row_index, row in enumerate(rows):
            for column_index, cell in enumerate(row):
                table.cell(row_index, column_index).text = _required_text(
                    cell,
                    allow_empty=True,
                )
        return
    raise DocumentEngineError(
        "document_block_unsupported",
        "Document block is unsupported",
    )


def _required_text(value: object, *, allow_empty: bool = False) -> str:
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        _invalid_content()
    return value


def _invalid_content() -> None:
    raise DocumentEngineError(
        "document_content_invalid",
        "Document content is invalid",
    )


def _normalize_ooxml(document: bytes) -> bytes:
    source = io.BytesIO(document)
    output = io.BytesIO()
    try:
        with zipfile.ZipFile(source) as archive, zipfile.ZipFile(
            output,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as normalized:
            for name in sorted(archive.namelist()):
                original = archive.getinfo(name)
                info = zipfile.ZipInfo(name, _FIXED_TIMESTAMP)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = original.external_attr
                normalized.writestr(info, archive.read(name))
    except (OSError, zipfile.BadZipFile) as error:
        raise DocumentEngineError(
            "document_create_failed",
            "DOCX creation failed",
        ) from error
    return output.getvalue()


def _verify_docx(document: bytes) -> None:
    if not document.startswith(b"PK"):
        _verification_failed("DOCX is not an OOXML package")
    required = {
        "[Content_Types].xml",
        "_rels/.rels",
        "word/document.xml",
        "word/styles.xml",
    }
    try:
        with zipfile.ZipFile(io.BytesIO(document)) as archive:
            names = archive.namelist()
            name_set = set(names)
            if len(name_set) != len(names):
                _verification_failed("DOCX contains duplicate package parts")
            if not required.issubset(name_set):
                _verification_failed("DOCX is missing required package parts")
            if archive.testzip() is not None:
                _verification_failed("DOCX package CRC validation failed")

            parsed: dict[str, ElementTree.Element] = {}
            for name in names:
                if name.endswith((".xml", ".rels")):
                    parsed[name] = ElementTree.fromstring(archive.read(name))
            _verify_relationships(parsed, name_set)
            text = "".join(
                node.text or ""
                for node in parsed["word/document.xml"].iter(
                    f"{{{_WORD_NAMESPACE}}}t"
                )
            )
            if not text.strip():
                _verification_failed("DOCX main document is empty")
    except DocumentEngineError:
        raise
    except (OSError, zipfile.BadZipFile, ElementTree.ParseError) as error:
        raise DocumentEngineError(
            "document_verification_failed",
            "DOCX package or XML is invalid",
        ) from error
    return None


def _verify_relationships(
    parsed: dict[str, ElementTree.Element],
    names: set[str],
) -> None:
    for relationship_path, root in parsed.items():
        if not relationship_path.endswith(".rels"):
            continue
        base = _relationship_base(relationship_path)
        for relationship in root.findall(
            f"{{{_RELATIONSHIP_NAMESPACE}}}Relationship"
        ):
            if relationship.get("TargetMode") == "External":
                continue
            target = relationship.get("Target")
            if not target:
                _verification_failed("DOCX relationship target is empty")
            resolved = posixpath.normpath(posixpath.join(base, target))
            if (
                resolved.startswith("../")
                or resolved.startswith("/")
                or resolved not in names
            ):
                _verification_failed(
                    f"DOCX relationship target is missing: {resolved}"
                )


def _relationship_base(relationship_path: str) -> str:
    path = PurePosixPath(relationship_path)
    if relationship_path == "_rels/.rels":
        return ""
    source_name = path.name.removesuffix(".rels")
    return str(path.parent.parent / source_name).rsplit("/", 1)[0]


def _verify_pdf(document: bytes) -> int:
    if not re.match(rb"%PDF-1\.[0-9]", document[:8]):
        _verification_failed("PDF header is invalid")
    if not re.search(rb"%%EOF\s*$", document):
        _verification_failed("PDF end marker is missing")

    start_matches = list(re.finditer(rb"startxref\s+(\d+)\s+%%EOF", document))
    if not start_matches:
        _verification_failed("PDF startxref is missing")
    xref_offset = int(start_matches[-1].group(1))
    if xref_offset < 0 or xref_offset >= len(document):
        _verification_failed("PDF cross-reference offset is invalid")
    xref_source = document[xref_offset:]
    if not (
        xref_source.startswith(b"xref")
        or re.match(
            rb"\d+\s+\d+\s+obj\b[\s\S]{0,2048}?/Type\s*/XRef\b",
            xref_source,
        )
    ):
        _verification_failed("PDF cross-reference offset is invalid")

    objects = {
        (int(match.group(1)), int(match.group(2))): match.group(3)
        for match in re.finditer(
            rb"(?m)^(\d+)\s+(\d+)\s+obj\b(.*?)\bendobj\b",
            document,
            re.DOTALL,
        )
    }
    if not objects:
        _verification_failed("PDF contains no indirect objects")

    root_reference = _pdf_reference(
        xref_source,
        rb"/Root\s+(\d+)\s+(\d+)\s+R",
        "PDF catalog reference is missing",
    )
    catalog = objects.get(root_reference)
    if catalog is None or not re.search(rb"/Type\s*/Catalog\b", catalog):
        _verification_failed("PDF catalog is invalid")
    pages_reference = _pdf_reference(
        catalog,
        rb"/Pages\s+(\d+)\s+(\d+)\s+R",
        "PDF page tree reference is missing",
    )
    pages = objects.get(pages_reference)
    if pages is None or not re.search(rb"/Type\s*/Pages\b", pages):
        _verification_failed("PDF page tree is invalid")
    count_match = re.search(rb"/Count\s+(\d+)\b", pages)
    if count_match is None or int(count_match.group(1)) <= 0:
        _verification_failed("PDF page count is invalid")
    page_count = int(count_match.group(1))
    concrete_pages = sum(
        1
        for body in objects.values()
        if re.search(rb"/Type\s*/Page\b", body)
    )
    if concrete_pages < page_count:
        _verification_failed("PDF page tree is incomplete")
    return page_count


def _pdf_reference(
    source: bytes,
    pattern: bytes,
    message: str,
) -> tuple[int, int]:
    match = re.search(pattern, source)
    if match is None:
        _verification_failed(message)
    return int(match.group(1)), int(match.group(2))


def _verification_failed(message: str) -> None:
    raise DocumentEngineError("document_verification_failed", message)
