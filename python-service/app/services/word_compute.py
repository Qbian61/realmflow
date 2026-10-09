from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


class WordComputeError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class WordComputeResult:
    document: bytes
    result: dict[str, object]
    modified: bool
    preservation_risk: list[dict[str, str]]


class WordComputeService:
    def compute(
        self,
        *,
        format_: str,
        operation: str,
        document: bytes,
        parameters: dict[str, object],
    ) -> WordComputeResult:
        if format_ != "docx":
            raise WordComputeError(
                "word_format_unsupported",
                "Word document format is unsupported",
            )
        risks = _detect_preservation_risk(document)
        try:
            word = Document(io.BytesIO(document))
        except Exception as error:
            raise WordComputeError(
                "word_invalid",
                "Word document is invalid",
            ) from error

        modified = operation not in {"inspect", "find"}
        if operation == "inspect":
            result = _inspect(word)
        elif operation == "find":
            result = _find(word, parameters)
        elif operation == "replace_text":
            result = _replace_text(word, parameters)
        elif operation == "insert_blocks":
            result = _insert_blocks(word, parameters)
        elif operation == "update_style":
            result = _update_style(word, parameters)
        elif operation == "update_layout":
            result = _update_layout(word, parameters)
        elif operation == "table_insert":
            result = _table_insert(word, parameters)
        elif operation == "table_write":
            result = _table_write(word, parameters)
        elif operation == "comment_add":
            result = _comment_add(word, parameters)
        elif operation == "comment_delete":
            result = _comment_delete(word, parameters)
        else:
            raise WordComputeError(
                "word_operation_unsupported",
                "Word document operation is unsupported",
            )

        output = document
        if modified:
            stream = io.BytesIO()
            word.save(stream)
            output = stream.getvalue()
        return WordComputeResult(
            document=output,
            result=result,
            modified=modified,
            preservation_risk=risks,
        )


def _inspect(document: DocumentObject) -> dict[str, object]:
    blocks: list[dict[str, object]] = []
    paragraph_index = 0
    table_index = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_index += 1
            blocks.append(_paragraph_block(block, paragraph_index))
        elif isinstance(block, Table):
            table_index += 1
            blocks.append(_table_block(block, table_index))
    return {
        "format": "docx",
        "blocks": blocks,
        "sections": [
            {
                "id": f"section-{index}",
                "widthPt": round(section.page_width.pt, 2),
                "heightPt": round(section.page_height.pt, 2),
                "orientation": (
                    "landscape"
                    if section.orientation == WD_ORIENT.LANDSCAPE
                    or section.page_width > section.page_height
                    else "portrait"
                ),
                "header": _story_text(section.header.paragraphs),
                "footer": _story_text(section.footer.paragraphs),
            }
            for index, section in enumerate(document.sections, start=1)
        ],
        "images": _images(document),
        "comments": _comments(document),
        "bookmarks": _bookmarks(document),
    }


def _paragraph_block(paragraph: Paragraph, index: int) -> dict[str, object]:
    style = paragraph.style.name if paragraph.style is not None else "Normal"
    heading = re.fullmatch(r"Heading ([1-9])", style, re.IGNORECASE)
    result: dict[str, object] = {
        "id": f"paragraph-{index}",
        "type": (
            "heading"
            if heading
            else "list_item"
            if style.lower().startswith("list")
            else "paragraph"
        ),
        "text": paragraph.text,
        "style": style,
    }
    if heading:
        result["level"] = int(heading.group(1))
    elif result["type"] == "list_item":
        result["level"] = _list_level(paragraph)
    return result


def _table_block(table: Table, index: int) -> dict[str, object]:
    return {
        "id": f"table-{index}",
        "type": "table",
        "rows": [
            [cell.text for cell in row.cells]
            for row in table.rows
        ],
    }


def _find(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    query = _require_string(parameters, "query")
    case_sensitive = parameters.get("caseSensitive") is True
    matches: list[dict[str, object]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        source = paragraph.text if case_sensitive else paragraph.text.casefold()
        needle = query if case_sensitive else query.casefold()
        start = 0
        while True:
            found = source.find(needle, start)
            if found < 0:
                break
            matches.append(
                {
                    "blockId": f"paragraph-{index}",
                    "start": found,
                    "end": found + len(query),
                    "text": paragraph.text[found : found + len(query)],
                }
            )
            start = found + max(1, len(query))
    return {"query": query, "matches": matches}


def _replace_text(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    query = _require_string(parameters, "query")
    replacement = _require_string(parameters, "replacement", allow_empty=True)
    replace_all = parameters.get("replaceAll") is True
    replaced = 0
    for paragraph in _all_paragraphs(document):
        while query in paragraph.text:
            _replace_first_in_runs(paragraph, query, replacement)
            replaced += 1
            if not replace_all:
                return {"replacements": replaced}
    if replaced == 0:
        raise WordComputeError(
            "word_text_not_found",
            "Text was not found in the Word document",
        )
    return {"replacements": replaced}


def _replace_first_in_runs(
    paragraph: Paragraph,
    query: str,
    replacement: str,
) -> None:
    full_text = "".join(run.text for run in paragraph.runs)
    start = full_text.find(query)
    if start < 0:
        return
    end = start + len(query)
    cursor = 0
    first_affected = None
    for index, run in enumerate(paragraph.runs):
        run_start = cursor
        run_end = cursor + len(run.text)
        cursor = run_end
        if run_end <= start or run_start >= end:
            continue
        if first_affected is None:
            first_affected = index
            prefix = run.text[: max(0, start - run_start)]
            suffix = run.text[max(0, end - run_start) :] if end <= run_end else ""
            run.text = f"{prefix}{replacement}{suffix}"
        else:
            suffix = run.text[max(0, end - run_start) :] if end < run_end else ""
            run.text = suffix


def _insert_blocks(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    markdown = _require_string(parameters, "markdown")
    after = _require_string(parameters, "afterBlockId")
    anchor = _require_top_level_block(document, after)
    inserted = _markdown_blocks(document, markdown)
    current = anchor._element
    for block in inserted:
        current.addnext(block._element)
        current = block._element
    return {"blocksInserted": len(inserted)}


def _markdown_blocks(
    document: DocumentObject,
    markdown: str,
) -> list[Paragraph | Table]:
    lines = markdown.splitlines()
    blocks: list[Paragraph | Table] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines):
            table_lines = [line]
            index += 1
            if _is_table_separator(lines[index]):
                index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [_split_table_row(value) for value in table_lines]
            table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = value
            blocks.append(table)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            paragraph = document.add_paragraph(
                style=f"Heading {len(heading.group(1))}"
            )
            paragraph.add_run(heading.group(2))
        elif re.match(r"^[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(re.sub(r"^[-*]\s+", "", line))
        else:
            paragraph = document.add_paragraph()
            _add_inline_markdown(paragraph, line)
        blocks.append(paragraph)
        index += 1
    return blocks


def _add_inline_markdown(paragraph: Paragraph, text: str) -> None:
    cursor = 0
    for match in re.finditer(r"\[([^\]]+)\]\((https?://[^)]+)\)", text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        _add_hyperlink(paragraph, match.group(1), match.group(2))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _add_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_style = OxmlElement("w:rStyle")
    run_style.set(qn("w:val"), "Hyperlink")
    run_properties.append(run_style)
    run.append(run_properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _update_style(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    paragraph = _require_paragraph(
        document,
        _require_string(parameters, "blockId"),
    )
    style = _require_mapping(parameters, "style")
    style_name = style.get("styleName")
    if style_name is not None:
        if not isinstance(style_name, str) or not style_name:
            raise _invalid_input("Word paragraph style is invalid")
        try:
            paragraph.style = style_name
        except KeyError as error:
            raise WordComputeError(
                "word_style_not_found",
                "Word paragraph style was not found",
            ) from error
    for run in paragraph.runs:
        if "bold" in style:
            run.bold = _require_bool(style, "bold")
        if "italic" in style:
            run.italic = _require_bool(style, "italic")
        if "fontSizePt" in style:
            run.font.size = Pt(_require_number(style, "fontSizePt"))
    return {"blockId": parameters["blockId"], "style": style}


def _update_layout(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    section = _require_section(
        document,
        _require_string(parameters, "sectionId"),
    )
    if "widthPt" in parameters:
        section.page_width = Pt(_require_number(parameters, "widthPt"))
    if "heightPt" in parameters:
        section.page_height = Pt(_require_number(parameters, "heightPt"))
    section.orientation = (
        WD_ORIENT.LANDSCAPE
        if section.page_width > section.page_height
        else WD_ORIENT.PORTRAIT
    )
    if "header" in parameters:
        _set_story_text(
            section.header.paragraphs,
            _require_string(parameters, "header", allow_empty=True),
        )
    if "footer" in parameters:
        _set_story_text(
            section.footer.paragraphs,
            _require_string(parameters, "footer", allow_empty=True),
        )
    return {"sectionId": parameters["sectionId"]}


def _table_insert(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    anchor = _require_top_level_block(
        document,
        _require_string(parameters, "afterBlockId"),
    )
    rows = _require_matrix(parameters, "rows")
    table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    anchor._element.addnext(table._element)
    return {"rows": len(rows), "columns": max(map(len, rows))}


def _table_write(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    table = _require_table(document, _require_string(parameters, "tableId"))
    start_row = _require_positive_integer(parameters, "startRow")
    start_column = _require_positive_integer(parameters, "startColumn")
    values = _require_matrix(parameters, "values")
    if (
        start_row + len(values) - 1 > len(table.rows)
        or any(
            start_column + len(row) - 1 > len(table.columns)
            for row in values
        )
    ):
        raise _invalid_input("Word table write exceeds the table bounds")
    for row_offset, row in enumerate(values):
        for column_offset, value in enumerate(row):
            table.cell(
                start_row - 1 + row_offset,
                start_column - 1 + column_offset,
            ).text = value
    return {"cellsWritten": sum(len(row) for row in values)}


def _comment_add(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    paragraph = _require_paragraph(
        document,
        _require_string(parameters, "blockId"),
    )
    if not paragraph.runs:
        paragraph.add_run("")
    comment = document.add_comment(
        paragraph.runs,
        text=_require_string(parameters, "text"),
        author=_require_string(parameters, "author"),
    )
    return {"commentId": str(comment.comment_id)}


def _comment_delete(
    document: DocumentObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    comment_id = _require_string(parameters, "commentId")
    try:
        numeric_id = int(comment_id)
    except ValueError as error:
        raise _invalid_input("Word comment id is invalid") from error
    comment = document.comments.get(numeric_id)
    if comment is None:
        raise WordComputeError(
            "word_comment_not_found",
            "Word comment was not found",
        )
    comments_element = document.comments._comments_elm
    comments_element.remove(comment._element)
    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        for node in document.element.body.xpath(
            f".//w:{tag}[@w:id='{numeric_id}']"
        ):
            node.getparent().remove(node)
    return {"commentId": comment_id, "deleted": True}


def _comments(document: DocumentObject) -> list[dict[str, object]]:
    paragraph_ids: dict[str, str] = {}
    for index, paragraph in enumerate(document.paragraphs, start=1):
        for node in paragraph._p.xpath(".//w:commentRangeStart"):
            paragraph_ids[node.get(qn("w:id"))] = f"paragraph-{index}"
    return [
        {
            "id": str(comment.comment_id),
            "author": comment.author,
            "text": comment.text,
            "blockId": paragraph_ids.get(
                str(comment.comment_id),
                "unknown",
            ),
        }
        for comment in document.comments
    ]


def _images(document: DocumentObject) -> list[dict[str, object]]:
    images: list[dict[str, object]] = []
    for index, shape in enumerate(document.inline_shapes, start=1):
        images.append(
            {
                "id": f"image-{index}",
                "widthPt": round(shape.width.pt, 2),
                "heightPt": round(shape.height.pt, 2),
            }
        )
    return images


def _bookmarks(document: DocumentObject) -> list[dict[str, str]]:
    return [
        {
            "id": node.get(qn("w:id")),
            "name": node.get(qn("w:name")),
        }
        for node in document.element.body.xpath(".//w:bookmarkStart")
        if node.get(qn("w:name"))
    ]


def _detect_preservation_risk(document: bytes) -> list[dict[str, str]]:
    definitions = (
        (
            b"<w:txbxContent",
            "word_text_box_unsupported",
            "Text boxes may not be preserved",
        ),
        (
            b"<w:fldSimple",
            "word_complex_field_unsupported",
            "Complex fields may not be preserved",
        ),
        (
            b"<w:instrText",
            "word_complex_field_unsupported",
            "Complex fields may not be preserved",
        ),
        (
            b"<w:ins",
            "word_tracked_changes_unsupported",
            "Tracked changes may not be preserved",
        ),
        (
            b"<w:del",
            "word_tracked_changes_unsupported",
            "Tracked changes may not be preserved",
        ),
        (
            b"<w:altChunk",
            "word_alt_chunk_unsupported",
            "Embedded alternative content may not be preserved",
        ),
    )
    risks: list[dict[str, str]] = []
    seen: set[str] = set()
    try:
        with zipfile.ZipFile(io.BytesIO(document)) as archive:
            for part in sorted(
                name
                for name in archive.namelist()
                if name.startswith("word/") and name.endswith(".xml")
            ):
                content = archive.read(part)
                for marker, code, message in definitions:
                    if marker in content and code not in seen:
                        risks.append(
                            {"code": code, "message": message, "part": part}
                        )
                        seen.add(code)
    except (zipfile.BadZipFile, KeyError) as error:
        raise WordComputeError(
            "word_invalid",
            "Word document is invalid",
        ) from error
    return risks


def _all_paragraphs(document: DocumentObject) -> Iterable[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _require_top_level_block(
    document: DocumentObject,
    block_id: str,
) -> Paragraph | Table:
    if block_id.startswith("paragraph-"):
        return _require_paragraph(document, block_id)
    if block_id.startswith("table-"):
        return _require_table(document, block_id)
    raise WordComputeError(
        "word_block_not_found",
        "Word document block was not found",
    )


def _require_paragraph(
    document: DocumentObject,
    block_id: str,
) -> Paragraph:
    index = _block_index(block_id, "paragraph")
    if index > len(document.paragraphs):
        raise WordComputeError(
            "word_block_not_found",
            "Word document paragraph was not found",
        )
    return document.paragraphs[index - 1]


def _require_table(document: DocumentObject, block_id: str) -> Table:
    index = _block_index(block_id, "table")
    if index > len(document.tables):
        raise WordComputeError(
            "word_block_not_found",
            "Word document table was not found",
        )
    return document.tables[index - 1]


def _require_section(document: DocumentObject, section_id: str) -> Any:
    index = _block_index(section_id, "section")
    if index > len(document.sections):
        raise WordComputeError(
            "word_section_not_found",
            "Word document section was not found",
        )
    return document.sections[index - 1]


def _block_index(value: str, prefix: str) -> int:
    match = re.fullmatch(rf"{prefix}-([1-9][0-9]*)", value)
    if not match:
        raise WordComputeError(
            "word_block_not_found",
            "Word document block was not found",
        )
    return int(match.group(1))


def _story_text(paragraphs: list[Paragraph]) -> str:
    return "\n".join(
        paragraph.text for paragraph in paragraphs if paragraph.text
    )


def _set_story_text(paragraphs: list[Paragraph], value: str) -> None:
    paragraph = paragraphs[0]
    paragraph.text = value
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def _list_level(paragraph: Paragraph) -> int:
    numbering = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    level = numbering.ilvl if numbering is not None else None
    return int(level.val) if level is not None else 0


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(
        re.fullmatch(r":?-{3,}:?", cell) is not None for cell in cells
    )


def _require_string(
    parameters: dict[str, object],
    key: str,
    *,
    allow_empty: bool = False,
) -> str:
    value = parameters.get(key)
    if not isinstance(value, str) or (not allow_empty and not value):
        raise _invalid_input(f"Word document {key} is invalid")
    return value


def _require_mapping(
    parameters: dict[str, object],
    key: str,
) -> dict[str, object]:
    value = parameters.get(key)
    if not isinstance(value, dict):
        raise _invalid_input(f"Word document {key} is invalid")
    return value


def _require_bool(parameters: dict[str, object], key: str) -> bool:
    value = parameters.get(key)
    if not isinstance(value, bool):
        raise _invalid_input(f"Word document {key} is invalid")
    return value


def _require_number(parameters: dict[str, object], key: str) -> float:
    value = parameters.get(key)
    if (
        not isinstance(value, (int, float))
        or isinstance(value, bool)
        or value <= 0
    ):
        raise _invalid_input(f"Word document {key} is invalid")
    return float(value)


def _require_positive_integer(
    parameters: dict[str, object],
    key: str,
) -> int:
    value = parameters.get(key)
    if not isinstance(value, int) or isinstance(value, bool) or value < 1:
        raise _invalid_input(f"Word document {key} is invalid")
    return value


def _require_matrix(
    parameters: dict[str, object],
    key: str,
) -> list[list[str]]:
    value = parameters.get(key)
    if (
        not isinstance(value, list)
        or not value
        or any(
            not isinstance(row, list)
            or not row
            or any(not isinstance(cell, str) for cell in row)
            for row in value
        )
    ):
        raise _invalid_input(f"Word document {key} is invalid")
    return value


def _invalid_input(message: str) -> WordComputeError:
    return WordComputeError("word_input_invalid", message)
