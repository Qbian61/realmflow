from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import range_boundaries


class SpreadsheetComputeError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class SpreadsheetComputeResult:
    document: bytes
    result: dict[str, object]
    modified: bool
    requires_recalculation: bool = False


class SpreadsheetComputeService:
    def compute(
        self,
        *,
        format_: str,
        operation: str,
        document: bytes,
        parameters: dict[str, object],
    ) -> SpreadsheetComputeResult:
        if format_ == "xlsx":
            return _compute_xlsx(operation, document, parameters)
        if format_ not in {"csv", "tsv"}:
            raise SpreadsheetComputeError(
                "spreadsheet_format_unsupported",
                "Spreadsheet format is unsupported",
            )
        delimiter = "," if format_ == "csv" else "\t"
        rows = _read_delimited(document, delimiter)
        if operation == "inspect":
            result = _inspect_delimited(rows, format_)
            modified = False
        elif operation == "read_range":
            result = _read_range(rows, parameters)
            modified = False
        elif operation == "write_range":
            result = _write_range(rows, parameters)
            modified = True
        elif operation == "insert_rows":
            result = _insert_rows(rows, parameters)
            modified = True
        elif operation == "delete_rows":
            result = _delete_rows(rows, parameters)
            modified = True
        elif operation == "sort":
            result = _sort_rows(rows, parameters)
            modified = True
        elif operation == "filter":
            result = _filter_rows(rows, parameters)
            modified = True
        else:
            raise SpreadsheetComputeError(
                "spreadsheet_operation_unsupported",
                "Spreadsheet operation is unsupported for this format",
            )
        return SpreadsheetComputeResult(
            document=_write_delimited(rows, delimiter),
            result=result,
            modified=modified,
        )


def _compute_xlsx(
    operation: str,
    document: bytes,
    parameters: dict[str, object],
) -> SpreadsheetComputeResult:
    try:
        workbook = load_workbook(io.BytesIO(document), data_only=False)
    except Exception as error:
        raise SpreadsheetComputeError(
            "spreadsheet_invalid",
            "Spreadsheet workbook is invalid",
        ) from error
    _assert_supported_workbook(workbook)
    modified = operation not in {"inspect", "read_range"}
    requires_recalculation = operation == "set_formula"
    if operation == "inspect":
        result = _inspect_xlsx(workbook)
    elif operation == "read_range":
        result = _read_xlsx_range(workbook, parameters)
    elif operation == "write_range":
        result = _write_xlsx_range(workbook, parameters, formulas=False)
    elif operation == "set_formula":
        result = _write_xlsx_range(workbook, parameters, formulas=True)
    elif operation == "set_style":
        result = _set_xlsx_style(workbook, parameters)
    elif operation == "insert_rows":
        result = _insert_xlsx_rows(workbook, parameters)
    elif operation == "delete_rows":
        result = _delete_xlsx_rows(workbook, parameters)
    elif operation == "sort":
        result = _sort_xlsx_rows(workbook, parameters)
    elif operation == "filter":
        result = _filter_xlsx(workbook, parameters)
    elif operation == "chart":
        result = _add_xlsx_chart(workbook, parameters)
    else:
        raise SpreadsheetComputeError(
            "spreadsheet_operation_unsupported",
            "Spreadsheet operation is unsupported for this format",
        )
    if not modified:
        output = document
    else:
        stream = io.BytesIO()
        workbook.save(stream)
        output = stream.getvalue()
    return SpreadsheetComputeResult(
        document=output,
        result=result,
        modified=modified,
        requires_recalculation=requires_recalculation,
    )


def _assert_supported_workbook(workbook: Any) -> None:
    if getattr(workbook, "vba_archive", None) is not None:
        raise SpreadsheetComputeError(
            "spreadsheet_macros_unsupported",
            "Macro-enabled workbooks are read-only",
        )
    if getattr(workbook, "_external_links", []):
        raise SpreadsheetComputeError(
            "spreadsheet_external_links_unsupported",
            "Workbooks with external links are read-only",
        )
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f" and not isinstance(cell.value, str):
                    raise SpreadsheetComputeError(
                        "spreadsheet_dynamic_formula_unsupported",
                        "Dynamic array formulas are read-only",
                    )


def _inspect_xlsx(workbook: Any) -> dict[str, object]:
    return {
        "format": "xlsx",
        "sheets": [
            {
                "name": sheet.title,
                "maxRow": sheet.max_row,
                "maxColumn": sheet.max_column,
                "usedRange": sheet.calculate_dimension(),
                "mergedRanges": [
                    str(merged) for merged in sheet.merged_cells.ranges
                ],
                "frozenPane": (
                    sheet.freeze_panes.coordinate
                    if hasattr(sheet.freeze_panes, "coordinate")
                    else sheet.freeze_panes
                ),
                "charts": [
                    {
                        "type": _chart_type(chart),
                        "title": _chart_title(chart),
                        "anchor": _chart_anchor(chart),
                    }
                    for chart in sheet._charts
                ],
            }
            for sheet in workbook.worksheets
        ],
        "hasFormulas": any(
            cell.data_type == "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        ),
        "hasExternalLinks": bool(getattr(workbook, "_external_links", [])),
        "hasMacros": getattr(workbook, "vba_archive", None) is not None,
    }


def _read_xlsx_range(
    workbook: Any,
    parameters: dict[str, object],
) -> dict[str, object]:
    sheet = _require_sheet(workbook, parameters)
    start_row, start_column, end_row, end_column = _require_range(parameters)
    rows: list[list[dict[str, object]]] = []
    for row in sheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_column,
        max_col=end_column,
    ):
        rows.append([_xlsx_cell(cell) for cell in row])
    return {
        "sheet": sheet.title,
        "range": _format_range(start_row, start_column, end_row, end_column),
        "rows": rows,
    }


def _xlsx_cell(cell: Any) -> dict[str, object]:
    value = None if isinstance(cell, MergedCell) else cell.value
    formula = value if cell.data_type == "f" and isinstance(value, str) else None
    result: dict[str, object] = {
        "address": cell.coordinate,
        "value": None if formula is not None else _json_cell_value(value),
        "valueType": _xlsx_value_type(value, formula),
    }
    if formula is not None:
        result["formula"] = formula
    if formula is not None or (value is not None and not isinstance(value, str)):
        result["numberFormat"] = cell.number_format
    if cell.has_style:
        result["style"] = {
            "font": {
                "bold": bool(cell.font.bold),
                "italic": bool(cell.font.italic),
                "color": _style_color(cell.font.color),
            },
            "fill": {"color": _style_color(cell.fill.fgColor)},
            "alignment": {
                "horizontal": cell.alignment.horizontal,
                "vertical": cell.alignment.vertical,
                "wrapText": bool(cell.alignment.wrap_text),
            },
        }
    return result


def _write_xlsx_range(
    workbook: Any,
    parameters: dict[str, object],
    *,
    formulas: bool,
) -> dict[str, object]:
    sheet = _require_sheet(workbook, parameters)
    start_row, start_column, end_row, end_column = _require_range(parameters)
    key = "formulas" if formulas else "values"
    values = parameters.get(key)
    if (
        not isinstance(values, list)
        or len(values) != end_row - start_row + 1
        or any(
            not isinstance(row, list) or len(row) != end_column - start_column + 1
            for row in values
        )
    ):
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            f"Spreadsheet {key} do not match the target range",
        )
    for row_offset, values_row in enumerate(values):
        for column_offset, value in enumerate(values_row):
            if formulas and (
                not isinstance(value, str) or not value.startswith("=")
            ):
                raise SpreadsheetComputeError(
                    "spreadsheet_input_invalid",
                    "Spreadsheet formula must start with equals",
                )
            if not formulas and not isinstance(
                value, (str, int, float, bool, type(None))
            ):
                raise SpreadsheetComputeError(
                    "spreadsheet_input_invalid",
                    "Spreadsheet cell value is invalid",
                )
            sheet.cell(
                row=start_row + row_offset,
                column=start_column + column_offset,
                value=value,
            )
    return {
        "sheet": sheet.title,
        "range": _format_range(start_row, start_column, end_row, end_column),
        "cellsWritten": (end_row - start_row + 1)
        * (end_column - start_column + 1),
    }


def _set_xlsx_style(
    workbook: Any,
    parameters: dict[str, object],
) -> dict[str, object]:
    sheet = _require_sheet(workbook, parameters)
    start_row, start_column, end_row, end_column = _require_range(parameters)
    style = parameters.get("style")
    if not isinstance(style, dict):
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            "Spreadsheet style is invalid",
        )
    font = style.get("font", {})
    fill = style.get("fill", {})
    alignment = style.get("alignment", {})
    if not all(isinstance(value, dict) for value in (font, fill, alignment)):
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            "Spreadsheet style is invalid",
        )
    for row in sheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_column,
        max_col=end_column,
    ):
        for cell in row:
            cell.font = Font(
                name=font.get("name", cell.font.name),
                size=font.get("size", cell.font.sz),
                bold=font.get("bold", cell.font.bold),
                italic=font.get("italic", cell.font.italic),
                color=font.get("color", _style_color(cell.font.color)),
            )
            if "color" in fill:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=str(fill["color"]),
                )
            cell.alignment = Alignment(
                horizontal=alignment.get(
                    "horizontal", cell.alignment.horizontal
                ),
                vertical=alignment.get("vertical", cell.alignment.vertical),
                wrap_text=alignment.get(
                    "wrapText", cell.alignment.wrap_text
                ),
            )
            if "numberFormat" in style:
                cell.number_format = str(style["numberFormat"])
    return {
        "sheet": sheet.title,
        "range": _format_range(start_row, start_column, end_row, end_column),
        "cellsStyled": (end_row - start_row + 1)
        * (end_column - start_column + 1),
    }


def _insert_xlsx_rows(
    workbook: Any,
    parameters: dict[str, object],
) -> dict[str, object]:
    sheet = _require_sheet(workbook, parameters)
    start_row = _positive_integer(parameters, "startRow")
    count = _positive_integer(parameters, "count")
    formulas = _capture_formulas(sheet)
    sheet.insert_rows(start_row, count)
    _shift_formulas(sheet, formulas, start_row, count)
    _shift_sheet_structures(sheet, start_row, count)
    return {"sheet": sheet.title, "startRow": start_row, "count": count}


def _delete_xlsx_rows(
    workbook: Any,
    parameters: dict[str, object],
) -> dict[str, object]:
    sheet = _require_sheet(workbook, parameters)
    start_row = _positive_integer(parameters, "startRow")
    count = _positive_integer(parameters, "count")
    if start_row + count - 1 > sheet.max_row:
        raise SpreadsheetComputeError(
            "spreadsheet_range_invalid",
            "Spreadsheet row is outside the used range",
        )
    sheet.delete_rows(start_row, count)
    _shift_sheet_structures(sheet, start_row + count, -count)
    return {"sheet": sheet.title, "startRow": start_row, "count": count}


def _sort_xlsx_rows(
    workbook: Any,
    parameters: dict[str, object],
) -> dict[str, object]:
    sheet = _require_sheet(workbook, parameters)
    start_row, start_column, end_row, end_column = _require_range(parameters)
    keys = parameters.get("keys")
    if not isinstance(keys, list) or not keys:
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            "Spreadsheet sort keys are required",
        )
    values = [
        [sheet.cell(row=row, column=column).value for column in range(start_column, end_column + 1)]
        for row in range(start_row, end_row + 1)
    ]
    for key in reversed(keys):
        if not isinstance(key, dict):
            raise SpreadsheetComputeError(
                "spreadsheet_input_invalid",
                "Spreadsheet sort key is invalid",
            )
        column = key.get("column")
        direction = key.get("direction")
        if (
            not isinstance(column, int)
            or column < start_column
            or column > end_column
            or direction not in {"ascending", "descending"}
        ):
            raise SpreadsheetComputeError(
                "spreadsheet_input_invalid",
                "Spreadsheet sort key is invalid",
            )
        offset = column - start_column
        values.sort(
            key=lambda row: _sort_value("" if row[offset] is None else str(row[offset])),
            reverse=direction == "descending",
        )
    for row_offset, row_values in enumerate(values):
        for column_offset, value in enumerate(row_values):
            sheet.cell(
                row=start_row + row_offset,
                column=start_column + column_offset,
                value=value,
            )
    return {
        "sheet": sheet.title,
        "range": _format_range(start_row, start_column, end_row, end_column),
        "rowsSorted": len(values),
    }


def _filter_xlsx(
    workbook: Any,
    parameters: dict[str, object],
) -> dict[str, object]:
    sheet = _require_sheet(workbook, parameters)
    start_row, start_column, end_row, end_column = _require_range(parameters)
    column = parameters.get("column")
    operator = parameters.get("operator")
    expected = parameters.get("value")
    if (
        not isinstance(column, int)
        or column < start_column
        or column > end_column
        or operator not in {"equals", "not_equals", "contains"}
        or not isinstance(expected, (str, int, float, bool))
    ):
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            "Spreadsheet filter is invalid",
        )
    matched = []
    for row in range(start_row + 1, end_row + 1):
        actual = sheet.cell(row=row, column=column).value
        if (
            (operator == "equals" and actual == expected)
            or (operator == "not_equals" and actual != expected)
            or (
                operator == "contains"
                and str(expected) in ("" if actual is None else str(actual))
            )
        ):
            matched.append(row)
    sheet.auto_filter.ref = _format_range(
        start_row, start_column, end_row, end_column
    )
    return {
        "sheet": sheet.title,
        "range": sheet.auto_filter.ref,
        "matchedRows": matched,
    }


def _add_xlsx_chart(
    workbook: Any,
    parameters: dict[str, object],
) -> dict[str, object]:
    sheet = _require_sheet(workbook, parameters)
    chart_type = parameters.get("type")
    chart_class = {
        "bar": BarChart,
        "line": LineChart,
        "pie": PieChart,
    }.get(chart_type)
    title = parameters.get("title")
    anchor = parameters.get("anchor")
    data_range = parameters.get("dataRange")
    category_range = parameters.get("categoryRange")
    if (
        chart_class is None
        or not isinstance(title, str)
        or not title
        or not isinstance(anchor, str)
        or not isinstance(data_range, str)
        or not isinstance(category_range, str)
    ):
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            "Spreadsheet chart definition is invalid",
        )
    data_bounds = range_boundaries(data_range)
    category_bounds = range_boundaries(category_range)
    chart = chart_class()
    chart.title = title
    chart.add_data(
        Reference(
            sheet,
            min_col=data_bounds[0],
            min_row=data_bounds[1],
            max_col=data_bounds[2],
            max_row=data_bounds[3],
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            sheet,
            min_col=category_bounds[0],
            min_row=category_bounds[1],
            max_col=category_bounds[2],
            max_row=category_bounds[3],
        )
    )
    sheet.add_chart(chart, anchor)
    return {
        "sheet": sheet.title,
        "type": chart_type,
        "title": title,
        "anchor": anchor,
    }


def _require_sheet(workbook: Any, parameters: dict[str, object]) -> Any:
    name = parameters.get("sheet")
    if not isinstance(name, str) or name not in workbook.sheetnames:
        raise SpreadsheetComputeError(
            "spreadsheet_sheet_not_found",
            "Spreadsheet sheet was not found",
        )
    return workbook[name]


def _capture_formulas(sheet: Any) -> list[tuple[int, int, str]]:
    return [
        (cell.row, cell.column, cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f" and isinstance(cell.value, str)
    ]


_CELL_REFERENCE = re.compile(r"(?<![A-Z0-9_])([$]?[A-Z]{1,3})([$]?)([1-9][0-9]*)")


def _shift_formula_references(formula: str, start_row: int, delta: int) -> str:
    def replace(match: re.Match[str]) -> str:
        row = int(match.group(3))
        if row < start_row:
            return match.group(0)
        return f"{match.group(1)}{match.group(2)}{max(1, row + delta)}"

    return _CELL_REFERENCE.sub(replace, formula)


def _shift_formulas(
    sheet: Any,
    formulas: list[tuple[int, int, str]],
    start_row: int,
    count: int,
) -> None:
    for old_row, column, formula in formulas:
        target_row = old_row + count if old_row >= start_row else old_row
        sheet.cell(
            row=target_row,
            column=column,
            value=_shift_formula_references(formula, start_row, count),
        )


def _shift_sheet_structures(sheet: Any, start_row: int, delta: int) -> None:
    if sheet.freeze_panes:
        coordinate = (
            sheet.freeze_panes.coordinate
            if hasattr(sheet.freeze_panes, "coordinate")
            else str(sheet.freeze_panes)
        )
        row, column = _parse_cell(coordinate)
        if row >= start_row:
            sheet.freeze_panes = f"{_column_name(column)}{max(1, row + delta)}"
    for chart in sheet._charts:
        anchor = chart.anchor
        if hasattr(anchor, "_from") and anchor._from.row >= start_row - 1:
            anchor._from.row = max(0, anchor._from.row + delta)
            if hasattr(anchor, "to"):
                anchor.to.row = max(anchor._from.row + 1, anchor.to.row + delta)
        for series in chart.ser:
            _shift_series_reference(series, start_row, delta)


def _shift_series_reference(series: Any, start_row: int, delta: int) -> None:
    for reference in (
        getattr(getattr(series, "val", None), "numRef", None),
        getattr(getattr(series, "cat", None), "numRef", None),
        getattr(getattr(series, "cat", None), "strRef", None),
    ):
        if reference is not None and isinstance(reference.f, str):
            reference.f = _shift_formula_references(
                reference.f, start_row, delta
            )


def _chart_type(chart: Any) -> str:
    name = chart.__class__.__name__.lower()
    return name.removesuffix("chart")


def _chart_title(chart: Any) -> str:
    try:
        paragraphs = chart.title.tx.rich.p
        return "".join(
            run.t
            for paragraph in paragraphs
            for run in paragraph.r
            if isinstance(run.t, str)
        )
    except (AttributeError, TypeError):
        return ""


def _chart_anchor(chart: Any) -> str:
    anchor = chart.anchor
    if isinstance(anchor, str):
        return anchor
    if hasattr(anchor, "_from"):
        return f"{_column_name(anchor._from.col + 1)}{anchor._from.row + 1}"
    return ""


def _json_cell_value(value: object) -> object:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _xlsx_value_type(value: object, formula: str | None) -> str:
    if formula is not None:
        return "formula"
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (datetime, date, time)):
        return "date"
    if isinstance(value, str):
        return "string"
    return "error"


def _style_color(color: Any) -> str | None:
    if color is None:
        return None
    value = getattr(color, "rgb", None)
    if isinstance(value, str):
        return value[-6:]
    return None


def _read_delimited(document: bytes, delimiter: str) -> list[list[str]]:
    try:
        text = document.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise SpreadsheetComputeError(
            "spreadsheet_encoding_unsupported",
            "Spreadsheet text must be UTF-8",
        ) from error
    try:
        return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    except csv.Error as error:
        raise SpreadsheetComputeError(
            "spreadsheet_invalid",
            "Spreadsheet text is invalid",
        ) from error


def _write_delimited(rows: list[list[str]], delimiter: str) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=delimiter, lineterminator="\n")
    writer.writerows(rows)
    return output.getvalue().encode("utf-8")


def _inspect_delimited(
    rows: list[list[str]],
    format_: str,
) -> dict[str, object]:
    max_column = max((len(row) for row in rows), default=0)
    used_range = (
        f"A1:{_column_name(max_column)}{len(rows)}"
        if rows and max_column
        else None
    )
    return {
        "format": format_,
        "sheets": [
            {
                "name": "Sheet1",
                "maxRow": len(rows),
                "maxColumn": max_column,
                "usedRange": used_range,
                "mergedRanges": [],
                "frozenPane": None,
                "charts": [],
            }
        ],
        "hasFormulas": any(
            value.startswith("=") for row in rows for value in row
        ),
        "hasExternalLinks": False,
        "hasMacros": False,
    }


def _read_range(
    rows: list[list[str]],
    parameters: dict[str, object],
) -> dict[str, object]:
    start_row, start_column, end_row, end_column = _require_range(parameters)
    result: list[list[dict[str, object]]] = []
    for row_number in range(start_row, end_row + 1):
        result_row: list[dict[str, object]] = []
        for column_number in range(start_column, end_column + 1):
            raw = _cell(rows, row_number, column_number)
            value, value_type, formula = _typed_value(raw)
            cell: dict[str, object] = {
                "address": f"{_column_name(column_number)}{row_number}",
                "value": value,
                "valueType": value_type,
            }
            if formula is not None:
                cell["formula"] = formula
            result_row.append(cell)
        result.append(result_row)
    return {
        "range": _format_range(start_row, start_column, end_row, end_column),
        "rows": result,
    }


def _write_range(
    rows: list[list[str]],
    parameters: dict[str, object],
) -> dict[str, object]:
    start_row, start_column, end_row, end_column = _require_range(parameters)
    values = parameters.get("values")
    if (
        not isinstance(values, list)
        or len(values) != end_row - start_row + 1
        or any(
            not isinstance(row, list) or len(row) != end_column - start_column + 1
            for row in values
        )
    ):
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            "Spreadsheet values do not match the target range",
        )
    _ensure_size(rows, end_row, end_column)
    for row_offset, values_row in enumerate(values):
        for column_offset, value in enumerate(values_row):
            rows[start_row - 1 + row_offset][
                start_column - 1 + column_offset
            ] = _serialize_value(value)
    return {
        "range": _format_range(start_row, start_column, end_row, end_column),
        "cellsWritten": (end_row - start_row + 1)
        * (end_column - start_column + 1),
    }


def _insert_rows(
    rows: list[list[str]],
    parameters: dict[str, object],
) -> dict[str, object]:
    start_row = _positive_integer(parameters, "startRow")
    count = _positive_integer(parameters, "count")
    if start_row > len(rows) + 1:
        raise SpreadsheetComputeError(
            "spreadsheet_range_invalid",
            "Spreadsheet row is outside the used range",
        )
    width = max((len(row) for row in rows), default=1)
    rows[start_row - 1 : start_row - 1] = [[""] * width for _ in range(count)]
    return {"startRow": start_row, "count": count}


def _delete_rows(
    rows: list[list[str]],
    parameters: dict[str, object],
) -> dict[str, object]:
    start_row = _positive_integer(parameters, "startRow")
    count = _positive_integer(parameters, "count")
    if start_row > len(rows) or start_row + count - 1 > len(rows):
        raise SpreadsheetComputeError(
            "spreadsheet_range_invalid",
            "Spreadsheet row is outside the used range",
        )
    del rows[start_row - 1 : start_row - 1 + count]
    return {"startRow": start_row, "count": count}


def _sort_rows(
    rows: list[list[str]],
    parameters: dict[str, object],
) -> dict[str, object]:
    start_row, start_column, end_row, end_column = _require_range(parameters)
    keys = parameters.get("keys")
    if not isinstance(keys, list) or not keys:
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            "Spreadsheet sort keys are required",
        )
    selected = rows[start_row - 1 : end_row]
    for key in reversed(keys):
        if not isinstance(key, dict):
            raise SpreadsheetComputeError(
                "spreadsheet_input_invalid",
                "Spreadsheet sort key is invalid",
            )
        column = key.get("column")
        direction = key.get("direction")
        if (
            not isinstance(column, int)
            or column < start_column
            or column > end_column
            or direction not in {"ascending", "descending"}
        ):
            raise SpreadsheetComputeError(
                "spreadsheet_input_invalid",
                "Spreadsheet sort key is invalid",
            )
        selected.sort(
            key=lambda row: _sort_value(
                row[column - 1] if column <= len(row) else ""
            ),
            reverse=direction == "descending",
        )
    rows[start_row - 1 : end_row] = selected
    return {
        "range": _format_range(start_row, start_column, end_row, end_column),
        "rowsSorted": len(selected),
    }


def _filter_rows(
    rows: list[list[str]],
    parameters: dict[str, object],
) -> dict[str, object]:
    start_row, start_column, end_row, end_column = _require_range(parameters)
    column = parameters.get("column")
    operator = parameters.get("operator")
    expected = parameters.get("value")
    if (
        not isinstance(column, int)
        or column < start_column
        or column > end_column
        or operator not in {"equals", "not_equals", "contains"}
        or not isinstance(expected, (str, int, float, bool))
    ):
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            "Spreadsheet filter is invalid",
        )
    expected_text = _serialize_value(expected)
    matched: list[int] = []
    for row_number in range(start_row + 1, end_row + 1):
        actual = _cell(rows, row_number, column)
        if (
            (operator == "equals" and actual == expected_text)
            or (operator == "not_equals" and actual != expected_text)
            or (operator == "contains" and expected_text in actual)
        ):
            matched.append(row_number)
    return {
        "range": _format_range(start_row, start_column, end_row, end_column),
        "matchedRows": matched,
    }


def _require_range(
    parameters: dict[str, object],
) -> tuple[int, int, int, int]:
    value = parameters.get("range")
    if not isinstance(value, str):
        raise SpreadsheetComputeError(
            "spreadsheet_range_invalid",
            "Spreadsheet range is invalid",
        )
    parts = value.upper().split(":")
    if len(parts) not in {1, 2}:
        raise SpreadsheetComputeError(
            "spreadsheet_range_invalid",
            "Spreadsheet range is invalid",
        )
    start = _parse_cell(parts[0])
    end = _parse_cell(parts[-1])
    if start[0] > end[0] or start[1] > end[1]:
        raise SpreadsheetComputeError(
            "spreadsheet_range_invalid",
            "Spreadsheet range is invalid",
        )
    return start[0], start[1], end[0], end[1]


def _parse_cell(value: str) -> tuple[int, int]:
    index = 0
    while index < len(value) and "A" <= value[index] <= "Z":
        index += 1
    letters, digits = value[:index], value[index:]
    if not letters or not digits or not digits.isdigit() or int(digits) < 1:
        raise SpreadsheetComputeError(
            "spreadsheet_range_invalid",
            "Spreadsheet range is invalid",
        )
    column = 0
    for letter in letters:
        column = column * 26 + ord(letter) - ord("A") + 1
    return int(digits), column


def _column_name(column: int) -> str:
    if column < 1:
        return "A"
    result = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def _format_range(
    start_row: int,
    start_column: int,
    end_row: int,
    end_column: int,
) -> str:
    return (
        f"{_column_name(start_column)}{start_row}:"
        f"{_column_name(end_column)}{end_row}"
    )


def _cell(rows: list[list[str]], row: int, column: int) -> str:
    if row > len(rows) or column > len(rows[row - 1]):
        return ""
    return rows[row - 1][column - 1]


def _typed_value(raw: str) -> tuple[object, str, str | None]:
    if raw.startswith("="):
        return None, "formula", raw
    lowered = raw.lower()
    if lowered == "true":
        return True, "boolean", None
    if lowered == "false":
        return False, "boolean", None
    try:
        return int(raw), "number", None
    except ValueError:
        try:
            return float(raw), "number", None
        except ValueError:
            pass
    try:
        if len(raw) == 10:
            date.fromisoformat(raw)
            return raw, "date", None
    except ValueError:
        pass
    if raw == "":
        return None, "blank", None
    return raw, "string", None


def _serialize_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (str, int, float)):
        return str(value)
    raise SpreadsheetComputeError(
        "spreadsheet_input_invalid",
        "Spreadsheet cell value is invalid",
    )


def _positive_integer(parameters: dict[str, object], key: str) -> int:
    value = parameters.get(key)
    if not isinstance(value, int) or isinstance(value, bool) or value < 1:
        raise SpreadsheetComputeError(
            "spreadsheet_input_invalid",
            f"Spreadsheet {key} is invalid",
        )
    return value


def _ensure_size(rows: list[list[str]], row_count: int, column_count: int) -> None:
    while len(rows) < row_count:
        rows.append([])
    for row in rows:
        if len(row) < column_count:
            row.extend([""] * (column_count - len(row)))


def _sort_value(value: str) -> tuple[int, object]:
    try:
        return 0, float(value)
    except ValueError:
        return 1, value.casefold()
