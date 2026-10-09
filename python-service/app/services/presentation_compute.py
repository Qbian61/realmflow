from __future__ import annotations

import base64
import binascii
import copy
import hashlib
import io
import zipfile
from dataclasses import dataclass
from typing import Any, Iterable

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.presentation import Presentation as PresentationObject
from pptx.slide import Slide
from pptx.util import Emu, Pt


class PresentationComputeError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class PresentationComputeResult:
    document: bytes
    result: dict[str, object]
    modified: bool
    preservation_risk: list[dict[str, str]]


MUTATING_OPERATIONS = {
    "update_text",
    "replace_image",
    "table_write",
    "chart_write",
    "add_slide",
    "copy_slide",
    "delete_slide",
    "reorder_slide",
    "add_text",
    "add_image",
    "add_table",
    "add_chart",
    "reorder_shape",
    "update_size",
}


class PresentationComputeService:
    def compute(
        self,
        *,
        format_: str,
        operation: str,
        document: bytes,
        parameters: dict[str, object],
    ) -> PresentationComputeResult:
        if format_ != "pptx":
            raise PresentationComputeError(
                "presentation_format_unsupported",
                "Presentation format is unsupported",
            )
        risks = _detect_preservation_risk(document)
        try:
            presentation = Presentation(io.BytesIO(document))
        except Exception as error:
            raise PresentationComputeError(
                "presentation_invalid",
                "Presentation document is invalid",
            ) from error

        if operation == "inspect":
            result = _inspect(presentation)
        elif operation == "update_text":
            result = _update_text(presentation, parameters)
        elif operation == "replace_image":
            result = _replace_image(presentation, parameters)
        elif operation == "table_write":
            result = _table_write(presentation, parameters)
        elif operation == "chart_write":
            result = _chart_write(presentation, parameters)
        elif operation == "add_slide":
            result = _add_slide(presentation, parameters)
        elif operation == "copy_slide":
            result = _copy_slide(presentation, parameters)
        elif operation == "delete_slide":
            result = _delete_slide(presentation, parameters)
        elif operation == "reorder_slide":
            result = _reorder_slide(presentation, parameters)
        elif operation == "add_text":
            result = _add_text(presentation, parameters)
        elif operation == "add_image":
            result = _add_image(presentation, parameters)
        elif operation == "add_table":
            result = _add_table(presentation, parameters)
        elif operation == "add_chart":
            result = _add_chart(presentation, parameters)
        elif operation == "reorder_shape":
            result = _reorder_shape(presentation, parameters)
        elif operation == "update_size":
            result = _update_size(presentation, parameters)
        else:
            raise PresentationComputeError(
                "presentation_operation_unsupported",
                "Presentation operation is unsupported",
            )

        modified = operation in MUTATING_OPERATIONS
        output = document
        if modified:
            stream = io.BytesIO()
            presentation.save(stream)
            output = stream.getvalue()
            risks = _detect_preservation_risk(output)
        return PresentationComputeResult(
            document=output,
            result=result,
            modified=modified,
            preservation_risk=risks,
        )


def _inspect(presentation: PresentationObject) -> dict[str, object]:
    slide_ids = list(presentation.slides._sldIdLst)
    return {
        "format": "pptx",
        "size": {
            "widthPt": _points(presentation.slide_width),
            "heightPt": _points(presentation.slide_height),
        },
        "slides": [
            _inspect_slide(slide, int(slide_ids[index - 1].id), index)
            for index, slide in enumerate(presentation.slides, start=1)
        ],
    }


def _inspect_slide(slide: Slide, slide_id: int, index: int) -> dict[str, object]:
    return {
        "slideId": _slide_id(slide_id),
        "index": index,
        "notes": _notes(slide),
        "shapes": [
            _inspect_shape(shape, slide_id, z_index)
            for z_index, shape in enumerate(slide.shapes, start=1)
        ],
    }


def _inspect_shape(
    shape: Any,
    slide_id: int,
    z_index: int,
) -> dict[str, object]:
    result: dict[str, object] = {
        "shapeId": _shape_id(slide_id, shape.shape_id),
        "name": shape.name,
        "type": _shape_type(shape),
        "zIndex": z_index,
        "bounds": {
            "leftPt": _points(shape.left),
            "topPt": _points(shape.top),
            "widthPt": _points(shape.width),
            "heightPt": _points(shape.height),
        },
    }
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        image = shape.image
        result["image"] = {
            "contentType": image.content_type,
            "widthPx": image.size[0],
            "heightPx": image.size[1],
            "sha256": hashlib.sha256(image.blob).hexdigest(),
        }
    elif shape.has_table:
        result["table"] = {
            "rows": [
                [cell.text for cell in row.cells]
                for row in shape.table.rows
            ]
        }
    elif shape.has_chart:
        result["chart"] = _inspect_chart(shape.chart)
    elif shape.has_text_frame:
        result["text"] = {
            "value": shape.text,
            "paragraphs": [
                {
                    "text": paragraph.text,
                    "alignment": (
                        paragraph.alignment.name.lower()
                        if paragraph.alignment is not None
                        else None
                    ),
                    "level": paragraph.level,
                }
                for paragraph in shape.text_frame.paragraphs
            ],
        }
    return result


def _inspect_chart(chart: Any) -> dict[str, object]:
    plot = chart.plots[0] if len(chart.plots) > 0 else None
    categories = (
        [str(category.label) for category in plot.categories]
        if plot is not None and plot.categories is not None
        else []
    )
    return {
        "type": _chart_type_name(chart.chart_type),
        "title": (
            chart.chart_title.text_frame.text
            if chart.has_title
            else ""
        ),
        "categories": categories,
        "series": [
            {
                "name": series.name,
                "values": [
                    float(value) if value is not None else None
                    for value in series.values
                ],
            }
            for series in chart.series
        ],
    }


def _update_text(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    shape = _require_shape(presentation, _require_string(parameters, "shapeId"))
    if not shape.has_text_frame:
        _invalid_shape_type("text")
    text = _require_string(parameters, "text", allow_empty=True)
    shape.text_frame.text = text
    return {"shapeId": _require_string(parameters, "shapeId"), "text": text}


def _replace_image(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    shape_id = _require_string(parameters, "shapeId")
    slide, shape = _require_shape_with_slide(presentation, shape_id)
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        _invalid_shape_type("image")
    image = _decode_base64(parameters, "imageBase64")
    try:
        _, relationship_id = slide.part.get_or_add_image_part(
            io.BytesIO(image)
        )
    except Exception as error:
        raise PresentationComputeError(
            "presentation_image_invalid",
            "Presentation image is invalid",
        ) from error
    shape._element.blipFill.blip.rEmbed = relationship_id
    return {
        "shapeId": shape_id,
        "sha256": hashlib.sha256(image).hexdigest(),
    }


def _table_write(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    shape_id = _require_string(parameters, "shapeId")
    shape = _require_shape(presentation, shape_id)
    if not shape.has_table:
        _invalid_shape_type("table")
    values = _require_matrix(parameters, "values")
    start_row = _require_integer(parameters, "startRow", minimum=1) - 1
    start_column = _require_integer(
        parameters,
        "startColumn",
        minimum=1,
    ) - 1
    table = shape.table
    if (
        start_row + len(values) > len(table.rows)
        or any(
            start_column + len(row) > len(table.columns)
            for row in values
        )
    ):
        raise PresentationComputeError(
            "presentation_table_range_invalid",
            "Presentation table write exceeds the table bounds",
        )
    written = 0
    for row_index, row in enumerate(values, start=start_row):
        for column_index, value in enumerate(row, start=start_column):
            table.cell(row_index, column_index).text = str(value)
            written += 1
    return {"shapeId": shape_id, "cellsWritten": written}


def _chart_write(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    shape_id = _require_string(parameters, "shapeId")
    slide, shape = _require_shape_with_slide(presentation, shape_id)
    if not shape.has_chart:
        _invalid_shape_type("chart")
    chart_data = _chart_data(parameters)
    chart_type = _chart_type(_chart_type_name(shape.chart.chart_type))
    replacement = slide.shapes.add_chart(
        chart_type,
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        chart_data,
    )
    title = _require_string(parameters, "title", allow_empty=True)
    replacement.chart.has_title = bool(title)
    if title:
        replacement.chart.chart_title.text_frame.text = title
    replacement.name = shape.name
    replacement.element.nvGraphicFramePr.cNvPr.id = shape.shape_id
    parent = shape.element.getparent()
    index = parent.index(shape.element)
    parent.remove(replacement.element)
    parent.remove(shape.element)
    parent.insert(index, replacement.element)
    return {
        "shapeId": shape_id,
        "seriesWritten": len(chart_data._series),
    }


def _add_slide(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    notes = parameters.get("notes")
    if notes is not None:
        if not isinstance(notes, str):
            _invalid_parameter("notes")
        slide.notes_slide.notes_text_frame.text = notes
    _move_slide_after(
        presentation,
        slide,
        parameters.get("afterSlideId"),
    )
    return {"slideId": _id_for_slide(presentation, slide)}


def _copy_slide(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    source = _require_slide(
        presentation,
        _require_string(parameters, "slideId"),
    )
    target = presentation.slides.add_slide(presentation.slide_layouts[6])
    relationship_ids = _copy_slide_relationships(source, target)
    for shape in source.shapes:
        element = copy.deepcopy(shape.element)
        for node in element.iter():
            for attribute, value in list(node.attrib.items()):
                if value in relationship_ids:
                    node.set(attribute, relationship_ids[value])
        target.shapes._spTree.insert_element_before(element, "p:extLst")
    if _notes(source):
        target.notes_slide.notes_text_frame.text = _notes(source)
    _move_slide_after(
        presentation,
        target,
        parameters.get("afterSlideId"),
    )
    return {"slideId": _id_for_slide(presentation, target)}


def _delete_slide(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    slide_id = _require_string(parameters, "slideId")
    if len(presentation.slides) <= 1:
        raise PresentationComputeError(
            "presentation_last_slide",
            "A presentation must retain at least one slide",
        )
    index, slide_element = _require_slide_element(presentation, slide_id)
    presentation.slides._sldIdLst.remove(slide_element)
    presentation.part.drop_rel(slide_element.rId)
    return {"slideId": slide_id, "deletedIndex": index + 1}


def _reorder_slide(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    slide_id = _require_string(parameters, "slideId")
    index = _require_integer(
        parameters,
        "index",
        minimum=1,
        maximum=len(presentation.slides),
    )
    _, element = _require_slide_element(presentation, slide_id)
    presentation.slides._sldIdLst.remove(element)
    presentation.slides._sldIdLst.insert(index - 1, element)
    return {"slideId": slide_id, "index": index}


def _add_text(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    slide, slide_number = _require_slide_and_number(
        presentation,
        _require_string(parameters, "slideId"),
    )
    bounds = _require_bounds(parameters)
    shape = slide.shapes.add_textbox(*bounds)
    shape.text = _require_string(parameters, "text", allow_empty=True)
    return {"shapeId": _shape_id(slide_number, shape.shape_id)}


def _add_image(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    slide, slide_number = _require_slide_and_number(
        presentation,
        _require_string(parameters, "slideId"),
    )
    image = _decode_base64(parameters, "imageBase64")
    try:
        shape = slide.shapes.add_picture(io.BytesIO(image), *_require_bounds(parameters))
    except Exception as error:
        raise PresentationComputeError(
            "presentation_image_invalid",
            "Presentation image is invalid",
        ) from error
    return {"shapeId": _shape_id(slide_number, shape.shape_id)}


def _add_table(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    slide, slide_number = _require_slide_and_number(
        presentation,
        _require_string(parameters, "slideId"),
    )
    rows = _require_matrix(parameters, "rows")
    column_count = max(len(row) for row in rows)
    shape = slide.shapes.add_table(
        len(rows),
        column_count,
        *_require_bounds(parameters),
    )
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            shape.table.cell(row_index, column_index).text = str(value)
    return {"shapeId": _shape_id(slide_number, shape.shape_id)}


def _add_chart(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    slide, slide_number = _require_slide_and_number(
        presentation,
        _require_string(parameters, "slideId"),
    )
    chart_type = _chart_type(_require_string(parameters, "type"))
    shape = slide.shapes.add_chart(
        chart_type,
        *_require_bounds(parameters),
        _chart_data(parameters),
    )
    title = _require_string(parameters, "title", allow_empty=True)
    shape.chart.has_title = bool(title)
    if title:
        shape.chart.chart_title.text_frame.text = title
    return {"shapeId": _shape_id(slide_number, shape.shape_id)}


def _reorder_shape(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    shape_id = _require_string(parameters, "shapeId")
    slide, shape = _require_shape_with_slide(presentation, shape_id)
    z_index = _require_integer(
        parameters,
        "zIndex",
        minimum=1,
        maximum=len(slide.shapes),
    )
    element = shape.element
    slide.shapes._spTree.remove(element)
    slide.shapes._spTree.insert(z_index + 1, element)
    return {"shapeId": shape_id, "zIndex": z_index}


def _update_size(
    presentation: PresentationObject,
    parameters: dict[str, object],
) -> dict[str, object]:
    width = _require_number(parameters, "widthPt", minimum=72)
    height = _require_number(parameters, "heightPt", minimum=72)
    presentation.slide_width = Pt(width)
    presentation.slide_height = Pt(height)
    return {"widthPt": width, "heightPt": height}


def _require_slide(
    presentation: PresentationObject,
    slide_id: str,
) -> Slide:
    slide, _ = _require_slide_and_number(presentation, slide_id)
    return slide


def _require_slide_and_number(
    presentation: PresentationObject,
    slide_id: str,
) -> tuple[Slide, int]:
    for element, slide in zip(
        presentation.slides._sldIdLst,
        presentation.slides,
        strict=True,
    ):
        numeric_id = int(element.id)
        if _slide_id(numeric_id) == slide_id:
            return slide, numeric_id
    raise PresentationComputeError(
        "presentation_slide_not_found",
        "Presentation slide was not found",
    )


def _require_slide_element(
    presentation: PresentationObject,
    slide_id: str,
) -> tuple[int, Any]:
    for index, element in enumerate(presentation.slides._sldIdLst):
        if _slide_id(int(element.id)) == slide_id:
            return index, element
    raise PresentationComputeError(
        "presentation_slide_not_found",
        "Presentation slide was not found",
    )


def _require_shape(
    presentation: PresentationObject,
    shape_id: str,
) -> Any:
    _, shape = _require_shape_with_slide(presentation, shape_id)
    return shape


def _require_shape_with_slide(
    presentation: PresentationObject,
    shape_id: str,
) -> tuple[Slide, Any]:
    for element, slide in zip(
        presentation.slides._sldIdLst,
        presentation.slides,
        strict=True,
    ):
        slide_number = int(element.id)
        for shape in slide.shapes:
            if _shape_id(slide_number, shape.shape_id) == shape_id:
                return slide, shape
    raise PresentationComputeError(
        "presentation_shape_not_found",
        "Presentation shape was not found",
    )


def _id_for_slide(
    presentation: PresentationObject,
    target: Slide,
) -> str:
    for element, slide in zip(
        presentation.slides._sldIdLst,
        presentation.slides,
        strict=True,
    ):
        if slide is target:
            return _slide_id(int(element.id))
    raise PresentationComputeError(
        "presentation_compute_invalid",
        "Presentation slide ID could not be resolved",
    )


def _move_slide_after(
    presentation: PresentationObject,
    slide: Slide,
    after_slide_id: object,
) -> None:
    target_id = _id_for_slide(presentation, slide)
    _, target_element = _require_slide_element(presentation, target_id)
    presentation.slides._sldIdLst.remove(target_element)
    if after_slide_id is None:
        presentation.slides._sldIdLst.insert(0, target_element)
        return
    if not isinstance(after_slide_id, str):
        _invalid_parameter("afterSlideId")
    index, _ = _require_slide_element(presentation, after_slide_id)
    presentation.slides._sldIdLst.insert(index + 1, target_element)


def _copy_slide_relationships(
    source: Slide,
    target: Slide,
) -> dict[str, str]:
    result: dict[str, str] = {}
    skipped = {RT.NOTES_SLIDE, RT.SLIDE_LAYOUT}
    for relationship_id, relationship in source.part.rels.items():
        if relationship.reltype in skipped:
            continue
        if relationship.is_external:
            new_id = target.part.relate_to(
                relationship.target_ref,
                relationship.reltype,
                is_external=True,
            )
        else:
            new_id = target.part.relate_to(
                relationship.target_part,
                relationship.reltype,
            )
        result[relationship_id] = new_id
    return result


def _notes(slide: Slide) -> str:
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text.strip()


def _shape_type(shape: Any) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if shape.has_table:
        return "table"
    if shape.has_chart:
        return "chart"
    if shape.has_text_frame:
        return "text"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    return "other"


def _chart_data(parameters: dict[str, object]) -> ChartData:
    categories = parameters.get("categories")
    series = parameters.get("series")
    if (
        not isinstance(categories, list)
        or not all(isinstance(value, (str, int, float)) for value in categories)
        or not isinstance(series, list)
        or len(series) == 0
    ):
        _invalid_parameter("chart data")
    data = ChartData()
    data.categories = [str(value) for value in categories]
    for item in series:
        if (
            not isinstance(item, dict)
            or not isinstance(item.get("name"), str)
            or not isinstance(item.get("values"), list)
            or len(item["values"]) != len(categories)
            or not all(
                isinstance(value, (int, float)) and not isinstance(value, bool)
                for value in item["values"]
            )
        ):
            _invalid_parameter("series")
        data.add_series(item["name"], item["values"])
    return data


def _chart_type(value: str) -> XL_CHART_TYPE:
    chart_types = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE,
        "pie": XL_CHART_TYPE.PIE,
    }
    if value not in chart_types:
        raise PresentationComputeError(
            "presentation_chart_type_invalid",
            "Presentation chart type is unsupported",
        )
    return chart_types[value]


def _chart_type_name(value: XL_CHART_TYPE) -> str:
    names = {
        XL_CHART_TYPE.BAR_CLUSTERED: "bar",
        XL_CHART_TYPE.COLUMN_CLUSTERED: "column",
        XL_CHART_TYPE.LINE: "line",
        XL_CHART_TYPE.PIE: "pie",
    }
    return names.get(value, value.name.lower())


def _require_bounds(parameters: dict[str, object]) -> tuple[Emu, Emu, Emu, Emu]:
    bounds = parameters.get("bounds")
    if not isinstance(bounds, dict):
        _invalid_parameter("bounds")
    return (
        Pt(_require_number(bounds, "leftPt", minimum=0)),
        Pt(_require_number(bounds, "topPt", minimum=0)),
        Pt(_require_number(bounds, "widthPt", minimum=1)),
        Pt(_require_number(bounds, "heightPt", minimum=1)),
    )


def _require_matrix(
    parameters: dict[str, object],
    key: str,
) -> list[list[object]]:
    value = parameters.get(key)
    if (
        not isinstance(value, list)
        or len(value) == 0
        or not all(isinstance(row, list) and len(row) > 0 for row in value)
    ):
        _invalid_parameter(key)
    return value


def _require_string(
    parameters: dict[str, object],
    key: str,
    *,
    allow_empty: bool = False,
) -> str:
    value = parameters.get(key)
    if not isinstance(value, str) or (not allow_empty and not value):
        _invalid_parameter(key)
    return value


def _require_integer(
    parameters: dict[str, object],
    key: str,
    *,
    minimum: int,
    maximum: int | None = None,
) -> int:
    value = parameters.get(key)
    if (
        not isinstance(value, int)
        or isinstance(value, bool)
        or value < minimum
        or (maximum is not None and value > maximum)
    ):
        _invalid_parameter(key)
    return value


def _require_number(
    parameters: dict[str, object],
    key: str,
    *,
    minimum: float,
) -> float:
    value = parameters.get(key)
    if (
        not isinstance(value, (int, float))
        or isinstance(value, bool)
        or value < minimum
    ):
        _invalid_parameter(key)
    return float(value)


def _decode_base64(parameters: dict[str, object], key: str) -> bytes:
    value = _require_string(parameters, key)
    try:
        decoded = base64.b64decode(value, validate=True)
    except (binascii.Error, ValueError) as error:
        raise PresentationComputeError(
            "presentation_image_invalid",
            "Presentation image encoding is invalid",
        ) from error
    if not decoded:
        raise PresentationComputeError(
            "presentation_image_invalid",
            "Presentation image is empty",
        )
    return decoded


def _invalid_parameter(name: str) -> None:
    raise PresentationComputeError(
        "presentation_parameter_invalid",
        f"Presentation parameter is invalid: {name}",
    )


def _invalid_shape_type(expected: str) -> None:
    raise PresentationComputeError(
        "presentation_shape_type_invalid",
        f"Presentation shape is not a {expected} shape",
    )


def _slide_id(value: int) -> str:
    return f"slide-{value}"


def _shape_id(slide_id: int, shape_id: int) -> str:
    return f"shape-{slide_id}-{shape_id}"


def _points(value: int | Emu) -> float:
    return round(float(value) / 12_700, 2)


def _detect_preservation_risk(
    document: bytes,
) -> list[dict[str, str]]:
    definitions: list[tuple[bytes, str, str]] = [
        (
            b"<p:timing",
            "presentation_animation_unsupported",
            "Animations and slide timing cannot be safely preserved",
        ),
        (
            b"<dgm:relIds",
            "presentation_smartart_unsupported",
            "SmartArt cannot be safely preserved",
        ),
        (
            b"<p:oleObj",
            "presentation_embedded_object_unsupported",
            "Embedded objects cannot be safely preserved",
        ),
    ]
    risks: list[dict[str, str]] = []
    try:
        with zipfile.ZipFile(io.BytesIO(document)) as archive:
            for part in sorted(archive.namelist()):
                if not part.endswith(".xml"):
                    continue
                content = archive.read(part)
                for marker, code, message in definitions:
                    if marker in content:
                        risks.append(
                            {"code": code, "message": message, "part": part}
                        )
    except (OSError, zipfile.BadZipFile):
        return []
    return risks
