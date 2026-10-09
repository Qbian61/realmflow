import asyncio
import io
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from docx import Document

from app import create_app
from app.services.document_engine import (
    DocumentExportResult,
    DocumentEngine,
    DocumentEngineError,
    libreoffice_pdf_command,
)


CONTENT = {
    "title": "Delivery report",
    "blocks": [
        {"kind": "heading", "level": 1, "text": "Summary"},
        {"kind": "paragraph", "text": "Ready for review."},
        {"kind": "bullet_list", "items": ["Deterministic", "Verified"]},
        {
            "kind": "table",
            "rows": [["Artifact", "State"], ["DOCX", "Ready"]],
        },
    ],
}


def client() -> TestClient:
    return TestClient(
        create_app(auth_token="session-secret"),
        headers={"Authorization": "Bearer session-secret"},
    )


def test_document_delivery_routes_create_export_and_verify(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.api import routes

    pdf = minimal_pdf()

    async def export_pdf(_document: bytes) -> DocumentExportResult:
        return DocumentExportResult(
            document=pdf,
            backend="libreoffice",
            quality="print",
            warnings=(),
        )

    monkeypatch.setattr(routes.document_engine, "export_pdf", export_pdf)

    created = client().post(
        "/api/v1/documents/create",
        json={"document": CONTENT},
    )
    assert created.status_code == 200, created.text
    created_payload = created.json()
    assert created_payload["format"] == "docx"
    assert created_payload["byteSize"] > 0
    assert created_payload["checksum"].startswith("sha256:")
    assert created_payload["pageCount"] is None

    exported = client().post(
        "/api/v1/documents/export-pdf",
        json={"documentBase64": created_payload["documentBase64"]},
    )
    assert exported.status_code == 200, exported.text
    assert exported.json() == {
        "documentBase64": __import__("base64").b64encode(pdf).decode("ascii"),
        "format": "pdf",
        "byteSize": len(pdf),
        "checksum": (
            "sha256:"
            + __import__("hashlib").sha256(pdf).hexdigest()
        ),
        "pageCount": 1,
        "converter": "libreoffice",
        "quality": "print",
        "warnings": [],
    }

    verified = client().post(
        "/api/v1/artifacts/verify",
        json={
            "format": "pdf",
            "documentBase64": exported.json()["documentBase64"],
        },
    )
    assert verified.status_code == 200, verified.text
    assert verified.json() == {
        "valid": True,
        "format": "pdf",
        "byteSize": len(pdf),
        "checksum": exported.json()["checksum"],
        "pageCount": 1,
    }


def test_document_delivery_routes_map_stable_engine_errors() -> None:
    response = client().post(
        "/api/v1/documents/export-pdf",
        json={"documentBase64": "bm90LWRvY3g="},
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": {
            "code": "document_verification_failed",
            "message": "DOCX is not an OOXML package",
        }
    }


def test_create_docx_is_byte_deterministic_and_preserves_structure() -> None:
    engine = DocumentEngine()

    first = engine.create_docx(CONTENT)
    second = engine.create_docx(CONTENT)

    assert first == second
    document = Document(io.BytesIO(first))
    assert document.core_properties.title == "Delivery report"
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Delivery report",
        "Summary",
        "Ready for review.",
        "Deterministic",
        "Verified",
    ]
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["Artifact", "State"],
        ["DOCX", "Ready"],
    ]
    verification = engine.verify(document=first, format_="docx")
    assert verification.valid is True
    assert verification.format == "docx"
    assert verification.page_count is None
    assert verification.byte_size == len(first)
    assert verification.checksum.startswith("sha256:")


def test_create_docx_rejects_unknown_or_empty_blocks() -> None:
    engine = DocumentEngine()

    with pytest.raises(DocumentEngineError) as empty:
        engine.create_docx({"title": "", "blocks": []})
    assert empty.value.code == "document_content_invalid"

    with pytest.raises(DocumentEngineError) as unknown:
        engine.create_docx(
            {"title": "Report", "blocks": [{"kind": "image", "path": "x"}]}
        )
    assert unknown.value.code == "document_block_unsupported"


def test_docx_verification_rejects_broken_relationship_targets() -> None:
    valid = DocumentEngine().create_docx(CONTENT)
    source = zipfile.ZipFile(io.BytesIO(valid))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"styles.xml", b"missing.xml")
            target.writestr(info, data)

    with pytest.raises(DocumentEngineError) as raised:
        DocumentEngine().verify(document=output.getvalue(), format_="docx")

    assert raised.value.code == "document_verification_failed"
    assert "relationship target" in str(raised.value)


def test_export_pdf_uses_isolated_libreoffice_and_deeply_verifies_output() -> None:
    observed: dict[str, Path | str] = {}
    pdf = minimal_pdf()

    async def run(
        executable: str,
        source_path: Path,
        output_path: Path,
        profile_path: Path,
    ) -> None:
        observed.update(
            executable=executable,
            source_path=source_path,
            output_path=output_path,
            profile_path=profile_path,
        )
        assert source_path.name == "source.docx"
        assert output_path.name == "source.pdf"
        assert profile_path.parent == source_path.parent
        output_path.write_bytes(pdf)

    engine = DocumentEngine(
        resolve_executable=lambda: "/usr/bin/soffice",
        run=run,
    )
    result = asyncio.run(engine.export_pdf(engine.create_docx(CONTENT)))

    assert result.document == pdf
    assert result.backend == "libreoffice"
    assert result.quality == "print"
    assert result.warnings == ()
    assert observed["executable"] == "/usr/bin/soffice"
    verification = engine.verify(document=result.document, format_="pdf")
    assert verification.valid is True
    assert verification.page_count == 1


def test_export_pdf_uses_cupsfilter_text_fallback_when_libreoffice_is_unavailable() -> None:
    source = DocumentEngine().create_docx(CONTENT)
    pdf = minimal_pdf()
    observed: dict[str, bytes | Path | str] = {}

    async def run_cupsfilter(
        executable: str,
        source_path: Path,
        output_path: Path,
    ) -> None:
        observed.update(
            executable=executable,
            source=source_path.read_bytes(),
            output_path=output_path,
        )
        output_path.write_bytes(pdf)

    engine = DocumentEngine(
        resolve_executable=lambda: None,
        resolve_cupsfilter=lambda: "/usr/sbin/cupsfilter",
        run_cupsfilter=run_cupsfilter,
    )

    result = asyncio.run(engine.export_pdf(source))

    assert result.document == pdf
    assert result.backend == "cupsfilter_text"
    assert result.quality == "degraded_text"
    assert result.warnings == (
        "PDF was generated from extracted plain text; original DOCX layout was not preserved.",
    )
    assert observed["executable"] == "/usr/sbin/cupsfilter"
    assert b"Delivery report" in observed["source"]
    assert b"Ready for review." in observed["source"]


def test_export_pdf_reports_unavailable_and_rejects_invalid_output() -> None:
    source = DocumentEngine().create_docx(CONTENT)
    unavailable = DocumentEngine(
        resolve_executable=lambda: None,
        resolve_cupsfilter=lambda: None,
    )

    with pytest.raises(DocumentEngineError) as missing:
        asyncio.run(unavailable.export_pdf(source))
    assert missing.value.code == "document_export_unavailable"

    async def bad_output(
        _executable: str,
        _source_path: Path,
        output_path: Path,
        _profile_path: Path,
    ) -> None:
        output_path.write_bytes(b"%PDF-1.7\ntruncated")

    invalid = DocumentEngine(
        resolve_executable=lambda: "/usr/bin/soffice",
        run=bad_output,
    )
    with pytest.raises(DocumentEngineError) as raised:
        asyncio.run(invalid.export_pdf(source))
    assert raised.value.code == "document_verification_failed"


def test_pdf_verification_rejects_invalid_cross_reference_offset() -> None:
    damaged = minimal_pdf().replace(b"startxref\n251", b"startxref\n999")

    with pytest.raises(DocumentEngineError) as raised:
        DocumentEngine().verify(document=damaged, format_="pdf")

    assert raised.value.code == "document_verification_failed"
    assert "cross-reference" in str(raised.value)


def test_libreoffice_pdf_command_is_fixed_and_argument_safe() -> None:
    command = libreoffice_pdf_command(
        executable="/usr/bin/soffice",
        source_path=Path("/tmp/source.docx"),
        output_path=Path("/tmp/output/source.pdf"),
        profile_path=Path("/tmp/profile"),
    )

    assert command == (
        "/usr/bin/soffice",
        "--headless",
        "--nologo",
        "--nodefault",
        "--nofirststartwizard",
        "-env:UserInstallation=file:///tmp/profile",
        "--convert-to",
        "pdf:writer_pdf_Export",
        "--outdir",
        "/tmp/output",
        "/tmp/source.docx",
    )


def minimal_pdf() -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R >>"
        ),
        b"<< /Length 0 >>\nstream\n\nendstream",
    ]
    output = bytearray(b"%PDF-1.7\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)
