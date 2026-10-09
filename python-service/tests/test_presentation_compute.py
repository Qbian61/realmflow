import base64
import hashlib
import io
import zipfile

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app import create_app


def client() -> TestClient:
    return TestClient(
        create_app(auth_token="session-secret"),
        headers={"Authorization": "Bearer session-secret"},
    )


def compute(
    operation: str,
    parameters: dict[str, object] | None = None,
    document: bytes | None = None,
) -> dict[str, object]:
    response = client().post(
        "/api/v1/office/presentations/compute",
        json={
            "format": "pptx",
            "operation": operation,
            "documentBase64": base64.b64encode(
                document or pptx_fixture()
            ).decode("ascii"),
            "parameters": parameters or {},
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def decoded(payload: dict[str, object]) -> bytes:
    return base64.b64decode(str(payload["documentBase64"]))


def test_inspects_ordered_slides_shapes_notes_and_boundaries() -> None:
    payload = compute("inspect")

    assert payload["modified"] is False
    assert payload["preservationRisk"] == []
    result = payload["result"]
    assert result["format"] == "pptx"
    assert result["size"] == {"widthPt": 720.0, "heightPt": 540.0}
    assert [slide["slideId"] for slide in result["slides"]] == [
        "slide-256",
        "slide-257",
    ]
    first = result["slides"][0]
    assert first["index"] == 1
    assert first["notes"] == "Explain the architecture"
    assert [shape["type"] for shape in first["shapes"]] == [
        "text",
        "image",
        "table",
        "chart",
    ]
    assert first["shapes"][0] == {
        "shapeId": "shape-256-2",
        "name": "Title",
        "type": "text",
        "zIndex": 1,
        "bounds": {
            "leftPt": 72.0,
            "topPt": 36.0,
            "widthPt": 288.0,
            "heightPt": 72.0,
        },
        "text": {
            "value": "RealmFlow",
            "paragraphs": [
                {
                    "text": "RealmFlow",
                    "alignment": None,
                    "level": 0,
                }
            ],
        },
    }
    assert first["shapes"][1]["image"] == {
        "contentType": "image/png",
        "widthPx": 8,
        "heightPx": 6,
        "sha256": hashlib.sha256(png_fixture("red")).hexdigest(),
    }
    assert first["shapes"][2]["table"]["rows"] == [
        ["Capability", "State"],
        ["PPTX", "Ready"],
    ]
    assert first["shapes"][3]["chart"] == {
        "type": "column",
        "title": "Progress",
        "categories": ["Read", "Write"],
        "series": [{"name": "Coverage", "values": [80.0, 60.0]}],
    }


def test_updates_text_image_table_and_chart_by_stable_shape_id() -> None:
    original = compute("inspect")["result"]["slides"][0]["shapes"]
    ids = {shape["type"]: shape["shapeId"] for shape in original}

    text = compute(
        "update_text",
        {"shapeId": ids["text"], "text": "Local-first"},
    )
    image = compute(
        "replace_image",
        {
            "shapeId": ids["image"],
            "imageBase64": base64.b64encode(png_fixture("blue")).decode("ascii"),
        },
        decoded(text),
    )
    table = compute(
        "table_write",
        {
            "shapeId": ids["table"],
            "startRow": 2,
            "startColumn": 2,
            "values": [["Complete"]],
        },
        decoded(image),
    )
    chart = compute(
        "chart_write",
        {
            "shapeId": ids["chart"],
            "title": "Delivery",
            "categories": ["Read", "Write"],
            "series": [{"name": "Coverage", "values": [100, 100]}],
        },
        decoded(table),
    )
    inspected = compute("inspect", document=decoded(chart))
    shapes = inspected["result"]["slides"][0]["shapes"]

    assert [shape["shapeId"] for shape in shapes] == list(ids.values())
    assert shapes[0]["text"]["value"] == "Local-first"
    assert shapes[1]["image"]["sha256"] == hashlib.sha256(
        png_fixture("blue")
    ).hexdigest()
    assert shapes[2]["table"]["rows"][1][1] == "Complete"
    assert shapes[3]["chart"]["title"] == "Delivery"
    assert shapes[3]["chart"]["series"][0]["values"] == [100.0, 100.0]


def test_adds_copies_reorders_and_deletes_slides_with_unique_ids() -> None:
    added = compute(
        "add_slide",
        {"afterSlideId": "slide-256", "notes": "Generated locally"},
    )
    added_inspection = compute("inspect", document=decoded(added))
    added_ids = [
        slide["slideId"] for slide in added_inspection["result"]["slides"]
    ]
    assert len(set(added_ids)) == 3
    new_id = added["result"]["slideId"]
    assert added_ids == ["slide-256", new_id, "slide-257"]

    copied = compute(
        "copy_slide",
        {"slideId": "slide-257", "afterSlideId": new_id},
        decoded(added),
    )
    copy_id = copied["result"]["slideId"]
    assert copy_id not in added_ids
    reordered = compute(
        "reorder_slide",
        {"slideId": copy_id, "index": 1},
        decoded(copied),
    )
    deleted = compute(
        "delete_slide",
        {"slideId": new_id},
        decoded(reordered),
    )
    final = compute("inspect", document=decoded(deleted))
    final_ids = [slide["slideId"] for slide in final["result"]["slides"]]
    assert final_ids == [copy_id, "slide-256", "slide-257"]


def test_editing_a_chart_on_a_copied_slide_does_not_modify_the_source() -> None:
    copied = compute(
        "copy_slide",
        {"slideId": "slide-256", "afterSlideId": "slide-256"},
    )
    copy_id = copied["result"]["slideId"]
    copied_inspection = compute("inspect", document=decoded(copied))
    copied_chart = next(
        shape
        for slide in copied_inspection["result"]["slides"]
        if slide["slideId"] == copy_id
        for shape in slide["shapes"]
        if shape["type"] == "chart"
    )

    changed = compute(
        "chart_write",
        {
            "shapeId": copied_chart["shapeId"],
            "title": "Copy only",
            "categories": ["Read", "Write"],
            "series": [{"name": "Coverage", "values": [10, 20]}],
        },
        decoded(copied),
    )
    final = compute("inspect", document=decoded(changed))
    charts = [
        next(shape for shape in slide["shapes"] if shape["type"] == "chart")
        for slide in final["result"]["slides"][:2]
    ]

    assert charts[0]["chart"]["title"] == "Progress"
    assert charts[1]["chart"]["title"] == "Copy only"


def test_adds_shapes_reorders_layers_and_updates_page_size() -> None:
    text = compute(
        "add_text",
        {
            "slideId": "slide-257",
            "text": "Summary",
            "bounds": {
                "leftPt": 36,
                "topPt": 36,
                "widthPt": 180,
                "heightPt": 54,
            },
        },
    )
    shape_id = text["result"]["shapeId"]
    table = compute(
        "add_table",
        {
            "slideId": "slide-257",
            "rows": [["Owner", "State"], ["RealmFlow", "Ready"]],
            "bounds": {
                "leftPt": 36,
                "topPt": 108,
                "widthPt": 288,
                "heightPt": 108,
            },
        },
        decoded(text),
    )
    chart = compute(
        "add_chart",
        {
            "slideId": "slide-257",
            "type": "bar",
            "title": "Status",
            "categories": ["Ready"],
            "series": [{"name": "Items", "values": [1]}],
            "bounds": {
                "leftPt": 360,
                "topPt": 108,
                "widthPt": 288,
                "heightPt": 216,
            },
        },
        decoded(table),
    )
    layered = compute(
        "reorder_shape",
        {"shapeId": shape_id, "zIndex": 3},
        decoded(chart),
    )
    sized = compute(
        "update_size",
        {"widthPt": 960, "heightPt": 540},
        decoded(layered),
    )
    inspected = compute("inspect", document=decoded(sized))

    assert inspected["result"]["size"] == {
        "widthPt": 960.0,
        "heightPt": 540.0,
    }
    second_shapes = inspected["result"]["slides"][1]["shapes"]
    assert [shape["type"] for shape in second_shapes] == [
        "table",
        "chart",
        "text",
    ]
    assert [shape["zIndex"] for shape in second_shapes] == [1, 2, 3]


def test_reports_animation_as_preservation_risk() -> None:
    inspected = compute("inspect", document=with_animation_risk(pptx_fixture()))

    assert inspected["preservationRisk"] == [
        {
            "code": "presentation_animation_unsupported",
            "message": "Animations and slide timing cannot be safely preserved",
            "part": "ppt/slides/slide1.xml",
        }
    ]


def test_rejects_invalid_document_encoding_with_stable_error() -> None:
    response = client().post(
        "/api/v1/office/presentations/compute",
        json={
            "format": "pptx",
            "operation": "inspect",
            "documentBase64": "***",
            "parameters": {},
        },
    )

    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "presentation_input_invalid"


def pptx_fixture() -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    title = slide.shapes.add_textbox(
        Inches(1),
        Inches(0.5),
        Inches(4),
        Inches(1),
    )
    title.name = "Title"
    title.text = "RealmFlow"
    slide.shapes.add_picture(
        io.BytesIO(png_fixture("red")),
        Inches(5.5),
        Inches(0.5),
        Inches(1),
        Inches(0.75),
    )
    table = slide.shapes.add_table(
        2,
        2,
        Inches(1),
        Inches(2),
        Inches(4),
        Inches(1.5),
    ).table
    table.cell(0, 0).text = "Capability"
    table.cell(0, 1).text = "State"
    table.cell(1, 0).text = "PPTX"
    table.cell(1, 1).text = "Ready"
    chart_data = ChartData()
    chart_data.categories = ["Read", "Write"]
    chart_data.add_series("Coverage", [80, 60])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(5.5),
        Inches(2),
        Inches(3.5),
        Inches(3),
        chart_data,
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Progress"
    slide.notes_slide.notes_text_frame.text = "Explain the architecture"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def png_fixture(color: str) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (8, 6), color=color).save(output, format="PNG")
    return output.getvalue()


def with_animation_risk(document: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(document))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "ppt/slides/slide1.xml":
                content = content.replace(
                    b"</p:sld>",
                    (
                        b"<p:timing><p:tnLst><p:par><p:cTn id=\"1\" "
                        b"dur=\"indefinite\"/></p:par></p:tnLst></p:timing>"
                        b"</p:sld>"
                    ),
                )
            target.writestr(item, content)
    return output.getvalue()
